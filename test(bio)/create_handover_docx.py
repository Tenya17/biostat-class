# -*- coding: utf-8 -*-
"""
生成 DM.xls 統計分析交接手冊 Word 文件 (.docx)
排版規範：
- 專業學術/研究團隊交接手冊風格
- 包含中文字體設定、表格美化、重點提醒區塊與圖表內嵌
"""

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """設定表格單元格背景顏色"""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    """設定單元格邊距 (單位: dxa, 20 dxa = 1 pt)"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_callout_box(doc, title, text, bg_color="F2F4F8", border_color="2B5C8F"):
    """建立類似 Alert / Note 的重點提醒方框"""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    
    cell = table.cell(0, 0)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=160, bottom=160, left=200, right=160)
    
    # 邊框設定：僅保留左側粗線條
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="36" w:space="0" w:color="{border_color}"/>'
        f'  <w:top w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run_title = p.add_run(f"【{title}】\n")
    run_title.font.name = '微軟正黑體'
    run_title.font.size = Pt(11)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x2B, 0x5C, 0x8F)
    
    run_text = p.add_run(text)
    run_text.font.name = '微軟正黑體'
    run_text.font.size = Pt(10)
    run_text.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    
    doc.add_paragraph() # 空行間距

def build_handover_document():
    doc = Document()
    
    # 版面邊界 (標準 1 英吋 = 2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
    
    # 預設字型
    doc.styles['Normal'].font.name = '微軟正黑體'
    doc.styles['Normal'].font.size = Pt(11)
    doc.styles['Normal'].font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.25
    doc.styles['Normal'].paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # 封面/文件標頭
    # -------------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(10)
    title_p.paragraph_format.space_after = Pt(4)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = title_p.add_run("DM.xls 生醫資料分析專案交接手冊")
    r_title.font.name = '微軟正黑體'
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(18)
    r_sub = sub_p.add_run("BMI、空腹血糖與收縮壓之統計方法解析、盒狀圖繪製與實務白話指南")
    r_sub.font.name = '微軟正黑體'
    r_sub.font.size = Pt(12)
    r_sub.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    
    # 中繼資訊表格 (元數據)
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.columns[0].width = Inches(1.8)
    meta_table.columns[1].width = Inches(4.7)
    
    meta_info = [
        ("專案檔案", "DM.xls (共 22 筆受試者資料)"),
        ("核心目標", "分析 BMI、血糖、收縮壓之組別差異，並以 Seaborn 繪製 Box plot"),
        ("交接對象", "接手研究員、實驗室同仁、臨床研究助理或非統計背景協同人員"),
        ("使用環境", "Python 3.9+ (搭配 pandas, scipy, seaborn, matplotlib)")
    ]
    
    for i, (k, v) in enumerate(meta_info):
        c0 = meta_table.cell(i, 0)
        c1 = meta_table.cell(i, 1)
        set_cell_background(c0, "F0F4F8")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(k)
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(v)
        r1.font.size = Pt(9.5)
        
    doc.add_paragraph() # 間距

    # -------------------------------------------------------------
    # 前言：交接背景與這份文件的目的
    # -------------------------------------------------------------
    h1 = doc.add_heading("一、 專案交接背景與目標", level=1)
    h1.style.font.name = '微軟正黑體'
    h1.style.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    
    p = doc.add_paragraph()
    p.add_run(
        "這份交接手冊主要記錄 DM.xls 資料集的分析過程。這批資料來自 22 位受試者，"
        "主要研究的是糖尿病患者（第 1 組）與對照健康組（第 2 組）在三項關鍵生理指標——"
    )
    r_bold = p.add_run("體位指標（BMI）、空腹血糖（Fasting Blood Glucose）與收縮壓（Systolic BP）")
    r_bold.font.bold = True
    p.add_run(
        "上的差異情形。為了確保未來的研究報告或論文投稿時，統計結果能經得起醫學同儕審查（Peer Review），"
        "我們在處理這份資料時，特別注重「統計方法的選擇邏輯」以及「繪圖呈現的真實度」。\n"
        "撰寫這份文件的目的，是讓接手的夥伴即使沒有深厚的統計背景，也能清楚理解「我們為什麼要這樣算」、"
        "「圖表該怎麼解讀」，以及「如何向醫師或主管用白話文解釋這些數字背後的意義」。"
    )

    # -------------------------------------------------------------
    # 第二章：專案檔案結構與重現方式
    # -------------------------------------------------------------
    h1 = doc.add_heading("二、 交付檔案清單與使用說明", level=1)
    h1.style.font.name = '微軟正黑體'
    h1.style.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p = doc.add_paragraph(
        "專案工作目錄下已準備好所有必要程式、匯出資料表與高畫質圖表，接手時請確認以下檔案皆齊全："
    )

    file_list = [
        ("DM.xls", "原始 Excel 資料檔，包含編號、組別 (1/2)、性別 (1/2)、年齡、收縮壓、血糖與 BMI。"),
        ("analysis.py", "主要 Python 分析程式。自動完成數據清洗、常態性與同質性檢驗、t 檢定、無母數檢定，並使用 Seaborn 繪製與匯出箱形圖。"),
        ("statistical_summary.csv", "分析結果匯出表，整理了全體與各組別的平均值、中位數、IQR、95% 信賴區間及所有檢定統計量。"),
        ("dm_boxplot_by_group.png", "【核心圖表】依組別區分之 BMI、血糖與收縮壓箱形圖（300 DPI 高畫質，疊加原始資料散點與 *** 顯著性標註）。"),
        ("dm_boxplot_overall.png", "全體樣本分佈箱形圖，標示了全體平均線與中位數線。"),
        ("dm_boxplot_gender.png", "組別 × 性別雙因子箱形圖，供次族群分佈探討。")
    ]
    
    t_files = doc.add_table(rows=len(file_list)+1, cols=2)
    t_files.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_files.columns[0].width = Inches(2.2)
    t_files.columns[1].width = Inches(4.3)
    
    # 標題列
    h_c0 = t_files.cell(0, 0)
    h_c1 = t_files.cell(0, 1)
    set_cell_background(h_c0, "2B5C8F")
    set_cell_background(h_c1, "2B5C8F")
    set_cell_margins(h_c0, 80, 80, 100, 100)
    set_cell_margins(h_c1, 80, 80, 100, 100)
    
    r = h_c0.paragraphs[0].add_run("檔案名稱")
    r.font.bold = True
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    r = h_c1.paragraphs[0].add_run("功能說明與用途")
    r.font.bold = True
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for idx, (fn, desc) in enumerate(file_list):
        row_cells = t_files.rows[idx+1].cells
        bg = "F9FBFC" if idx % 2 == 1 else "FFFFFF"
        set_cell_background(row_cells[0], bg)
        set_cell_background(row_cells[1], bg)
        set_cell_margins(row_cells[0], 60, 60, 100, 100)
        set_cell_margins(row_cells[1], 60, 60, 100, 100)
        
        p0 = row_cells[0].paragraphs[0]
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(fn)
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        
        p1 = row_cells[1].paragraphs[0]
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(desc)
        r1.font.size = Pt(9.5)

    doc.add_paragraph()

    # 執行說明
    p_run = doc.add_paragraph()
    p_run.add_run("【如何重新執行分析】\n").font.bold = True
    p_run.add_run(
        "只要在該目錄的終端機執行：\n"
        "    .venv/bin/python analysis.py\n"
        "程式會自動重新計算所有統計檢定、列印詳細報表至螢幕，並在數秒內重新生成所有圖檔與 CSV 摘要表。"
    )

    # -------------------------------------------------------------
    # 第三章：統計方法深度解析（交接重點：為什麼要這樣做？）
    # -------------------------------------------------------------
    h1 = doc.add_heading("三、 核心統計方法與選擇邏輯（重點章節）", level=1)
    h1.style.font.name = '微軟正黑體'
    h1.style.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p = doc.add_paragraph()
    p.add_run(
        "在醫學小樣本數據（本案每組只有 11 人）的分析中，"
    )
    p.add_run("「最忌諱的是直接盲目套用 t 檢定」").font.bold = True
    p.add_run(
        "。若沒有先驗證資料的先決條件（Assumptions），往往會被口試委員或期刊審稿人質疑結論的可靠性。\n"
        "因此，本次分析我們採取了「嚴謹三步驟檢驗原則」，依序執行常態性、變異數同質性檢定，"
        "並同時提供母數與無母數檢定雙軌並行，具體決策過程如下："
    )

    # 步驟 1
    h2 = doc.add_heading("步驟 1：常態性檢定 —— 為什麼選 Shapiro-Wilk 檢定？", level=2)
    h2.style.font.name = '微軟正黑體'
    
    p = doc.add_paragraph()
    p.add_run(
        "• 方法原理與選擇考量：\n"
        "  一般常見的常態檢定有 Kolmogorov-Smirnov (K-S) 與 Shapiro-Wilk。但統計學界公認：當樣本數少於 50 時，"
        "K-S 檢定的檢力極低，容易發生型二錯誤（明明不常態卻誤判為常態）。因此，國際生醫研究針對小樣本，一律推薦使用 "
    )
    p.add_run("Shapiro-Wilk 檢定").font.bold = True
    p.add_run("。\n")
    p.add_run(
        "• 白話文解讀：\n"
        "  常態分佈就像是標準的鐘形曲線。Shapiro-Wilk 檢定的虛無假說是「資料符合常態分佈」。\n"
        "  若 p 值 > 0.05，代表資料沒有明顯偏離常態；若 p 值 < 0.05，代表資料歪斜嚴重，不宜隨意使用常規平均值與標準差。\n"
        "• 本專案數據表現：\n"
        "  本案中，第 1 組與第 2 組的 BMI（p=0.284, 0.447）、血糖（p=0.461, 0.311）、收縮壓（p=0.699, 0.512）"
        "所有 p 值皆大於 0.05，這表示兩組在各項指標上皆大致符合常態分佈假設。"
    )

    # 步驟 2
    h2 = doc.add_heading("步驟 2：變異數同質性檢定 —— 為什麼血糖不能用傳統 t 檢定？", level=2)
    h2.style.font.name = '微軟正黑體'

    p = doc.add_paragraph()
    p.add_run(
        "• 方法原理與選擇考量：\n"
        "  使用獨立樣本 t 檢定比較兩組平均值時，有一個關鍵前提是「兩組資料的離散程度（變異數）必須差不多大（齊一性）」。"
        "我們採用 Levene's Test 來檢驗這個前提。\n"
        "• 白話文解讀：\n"
        "  就像比較兩支球隊的身高，如果一隊隊員大家身高都很平均（分散小），另一隊有人極高有人極矮（分散極大），"
        "我們就不能假定兩隊的分散程度相同。\n"
        "• 本專案數據表現（關鍵差異點！）：\n"
        "  - BMI (Levene p = 0.4530) 與收縮壓 (Levene p = 0.9437)：p > 0.05，變異數齊一，可直接使用標準 Student's t 檢定。\n"
        "  - "
    )
    r_alert = p.add_run("血糖 (Levene p = 0.0391 < 0.05)：")
    r_alert.font.bold = True
    r_alert.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
    p.add_run(
        "檢定結果顯著，代表兩組變異數不相等！第 1 組糖尿病患者血糖標準差高達 12.04，而健康組只有 6.28（患者組的血糖波動範圍是對照組的兩倍）。\n"
        "  因此在程式中，我們針對血糖自動啟用了 "
    )
    p.add_run("Welch's t-test（自由度校正）").font.bold = True
    p.add_run("，避免因為變異數不齊而造成型一錯誤（False Positive）。這點在後續寫論文或簡報時務必向評審交代，顯現分析的專業度！")

    # 步驟 3
    h2 = doc.add_heading("步驟 3：無母數檢定並行 —— 曼-惠特尼 U 檢定 (Mann-Whitney U Test)", level=2)
    h2.style.font.name = '微軟正黑體'

    p = doc.add_paragraph()
    p.add_run(
        "• 為什麼每組 11 人一定要補放無母數檢定？\n"
        "  雖然 Shapiro-Wilk 檢定顯示常態，但小樣本的統計檢力有限。為了讓結論萬無一失，醫學文獻標準做法是同時提供無母數檢定作為佐證。\n"
        "• 白話文解讀：\n"
        "  無母數檢定不看「實際數值差多少」，而是把所有人從小排到大，看「排名名次（Rank）」。"
        "它不怕極端值干擾，是最穩健的檢驗方式。\n"
        "• 本專案數據的神奇結果：\n"
        "  三項指標計算出來的 Mann-Whitney U 值全部都是 "
    )
    p.add_run("121.0 (p = 8.15 × 10⁻⁵ = 0.00008)").font.bold = True
    p.add_run(
        "。\n"
        "  這在統計上意味著什麼？因為兩組各 11 人，$11 \times 11 = 121$ 就是理論上的最高上限！\n"
        "  也就是說：第 1 組中數值最低的那位受試者，依然比第 2 組中數值最高的那位受試者還要高！兩組數值完全沒有任何重疊交集。"
    )

    # 步驟 4
    h2 = doc.add_heading("步驟 4：效應值 (Cohen's d) —— 為什麼只看 p 值還不夠？", level=2)
    h2.style.font.name = '微軟正黑體'

    p = doc.add_paragraph()
    p.add_run(
        "• 觀念釐清：\n"
        "  很多人以為 p < 0.001 就代表「差異非常巨大」，這在統計學上是不完全正確的。p 值只告訴我們「兩組的差異是不是運氣巧合」，"
        "而只有 "
    )
    p.add_run("效應值（Effect Size，如 Cohen's d）").font.bold = True
    p.add_run(
        " 才能真正量化「差距到底有多大」。\n"
        "• 判讀基準：一般統計學界認為 d = 0.2 為小效應，d = 0.5 為中等效應，d ≥ 0.8 為大效應（Large Effect）。\n"
        "• 本數據結果：\n"
        "  - BMI 的 Cohen's d = 6.53\n"
        "  - 血糖的 Cohen's d = 10.19\n"
        "  - 收縮壓的 Cohen's d = 5.02\n"
        "  數值全部遠大於 0.8，這說明兩組並非只有「統計上的顯著」，在「臨床實務上」也是極端懸殊的差距。"
    )

    add_callout_box(
        doc,
        "統計方法交接一句話口訣",
        "小樣本先看 Shapiro（確認常態）→ 再測 Levene（看變異數齊不齊，不齊用 Welch 校正）"
        "→ t 檢定配合 Mann-Whitney U 雙軌印證 → 最後補上 Cohen's d 說明臨床差異幅度。",
        bg_color="EBF5FB",
        border_color="2980B9"
    )

    # -------------------------------------------------------------
    # 第四章：統計分析結果匯總表與數據解讀
    # -------------------------------------------------------------
    h1 = doc.add_heading("四、 數據分析成果匯總與臨床解讀", level=1)
    h1.style.font.name = '微軟正黑體'
    h1.style.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p = doc.add_paragraph("以下為三項指標的完整統計數值總整表（已四捨五入至易讀位數）：")

    # 匯總表格
    summary_data = [
        ("指標項目", "全體樣本\n(N=22)", "第 1 組 (個案組)\n(n=11)", "第 2 組 (對照組)\n(n=11)", "平均差\n(95% CI)", "t 檢定值 (p 值)", "效應值\n(Cohen's d)"),
        ("BMI\n(kg/m²)", "27.01 ± 3.92\n中位 27.35\n[21.7 - 32.3]", "30.68 ± 0.99\n中位 30.50\n[29.5 - 32.3]", "23.34 ± 1.25\n中位 23.10\n[21.7 - 25.2]", "7.35\n[6.35, 8.35]", "t = 15.32\np < 0.001", "d = 6.53\n(極大)"),
        ("空腹血糖\n(mg/dL)", "144.41 ± 50.98\n中位 142.00\n[88.0 - 212.0]", "193.36 ± 12.04\n中位 192.00\n[178.0 - 212.0]", "95.45 ± 6.28\n中位 94.00\n[88.0 - 106.0]", "97.91\n[89.18, 106.64]", "t = 23.91\np < 0.001\n(Welch)", "d = 10.19\n(極大)"),
        ("收縮壓\n(mmHg)", "138.77 ± 15.18\n中位 139.50\n[118.0 - 162.0]", "152.64 ± 5.55\n中位 152.00\n[145.0 - 162.0]", "124.91 ± 5.49\n中位 124.00\n[118.0 - 134.0]", "27.73\n[22.82, 32.64]", "t = 11.78\np < 0.001", "d = 5.02\n(極大)")
    ]

    t_res = doc.add_table(rows=len(summary_data), cols=7)
    t_res.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Inches(1.1), Inches(1.0), Inches(1.0), Inches(1.0), Inches(0.9), Inches(1.0), Inches(0.8)]
    for r_idx, row in enumerate(summary_data):
        for c_idx, val in enumerate(row):
            cell = t_res.cell(r_idx, c_idx)
            cell.width = widths[c_idx]
            if r_idx == 0:
                set_cell_background(cell, "2B5C8F")
                set_cell_margins(cell, 80, 80, 60, 60)
                p_c = cell.paragraphs[0]
                p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p_c.add_run(val)
                r.font.bold = True
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            else:
                bg = "F4F7F9" if r_idx % 2 == 1 else "FFFFFF"
                set_cell_background(cell, bg)
                set_cell_margins(cell, 60, 60, 60, 60)
                p_c = cell.paragraphs[0]
                p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p_c.add_run(val)
                r.font.size = Pt(8.5)

    doc.add_paragraph()

    # 數據臨床口語化白話翻譯
    p_interp = doc.add_paragraph()
    p_interp.add_run("【一般人一聽就懂的白話結論】\n").font.bold = True
    p_interp.add_run(
        "1. 體位（BMI）：第 1 組平均 BMI 高達 30.68 kg/m²（醫學定義已達中度肥胖），而第 2 組為 23.34 kg/m²（完全在健康理想體位 18.5~24 範圍內）。"
        "兩組平均差了 7.35 個單位，在體態上有著截然不同的分界。\n"
        "2. 血糖：第 1 組平均血糖高達 193.36 mg/dL（臨床上空腹血糖 ≥ 126 即確診為糖尿病），而第 2 組為 95.45 mg/dL（完全在正常值 < 100 內）。"
        "兩組血糖差距近 100 mg/dL，病理特徵非常明確。\n"
        "3. 收縮壓：第 1 組平均收縮壓為 152.64 mmHg（達到第 2 期高血壓 ≥ 140 mmHg 標準），對照組為 124.91 mmHg。"
        "平均高出近 28 mmHg，心血管負擔顯著偏高。\n"
        "4. 指標連動性（相關係數 r > 0.98）：這三項數值並不是各自獨立的，而是高度「同進同退」——BMI 越高的受試者，其血糖與血壓幾乎百分之百同步飆高，"
        "非常典型地印證了「代謝症候群（Metabolic Syndrome）」三位一體的臨床特徵。"
    )

    # -------------------------------------------------------------
    # 第五章：Seaborn 箱形圖繪圖理念與圖表判讀手冊
    # -------------------------------------------------------------
    h1 = doc.add_heading("五、 箱形圖（Box Plot）視覺化理念與判讀手冊", level=1)
    h1.style.font.name = '微軟正黑體'
    h1.style.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p = doc.add_paragraph()
    p.add_run(
        "在資料呈現上，許多人習慣只畫傳統的長條圖（Bar chart）加誤差線。但在現代生物醫學領域中，長條圖容易隱藏真實資料分佈與極端值。\n"
        "我們依照交接指示，採用 Python "
    )
    p.add_run("Seaborn").font.bold = True
    p.add_run(
        " 繪製箱形圖（Box plot），並加入了「資料點抖動散點（Stripplot/Jitter）」技術，以下說明圖表的設計重點與閱讀方式："
    )

    # 盒狀圖圖解
    p_box_desc = doc.add_paragraph()
    p_box_desc.add_run("【如何向非專業人士解讀箱形圖？】\n").font.bold = True
    p_box_desc.add_run(
        "• 中間有顏色的箱子（Box）：代表這組人中「中間 50%」的資料範圍（從第 25 百分位數 Q1 到第 75 百分位數 Q3）。箱子越扁，代表大家數值越集中；箱子越高，代表分散度大。\n"
        "• 箱子中間的黑粗橫線：代表「中位數（Median）」，也就是剛好排在最中間那位病患的數值（比平均值更不容易被極端值拉偏）。\n"
        "• 上下延伸的黑線（鬚線 Whiskers）：代表正常範圍內的極大值與極小值（一般定義為 1.5 倍四分位距以內）。\n"
        "• 黑色散落圓點（Stripplot 點位）：這是每一位受試者的真實原始數據。透過透明度與微幅左右抖動（Jitter），讓讀者一眼看清「22 個人每個人落在哪裡」，完全公開透明，不隱藏任何原始資訊。\n"
        "• 菱形標記（Diamond）：代表該組的「平均值（Mean）」，方便同時比對中位數與平均數是否一致。\n"
        "• 頂部橫槓標註（*** p < 0.001）：代表經過統計檢定，兩組差異在統計學上達到極端顯著。"
    )

    # 嵌入核心圖表 1
    h2 = doc.add_heading("核心圖表一：兩組指標箱形圖比較（dm_boxplot_by_group.png）", level=2)
    h2.style.font.name = '微軟正黑體'
    
    img_group_path = "dm_boxplot_by_group.png"
    if os.path.exists(img_group_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_group_path, width=Inches(6.2))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(12)
        r_cap = p_cap.add_run("圖 1：第 1 組 (糖尿病組) 與第 2 組 (對照組) 之 BMI、血糖與收縮壓箱形圖比較")
        r_cap.font.size = Pt(9.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    # 嵌入圖表 2
    h2 = doc.add_heading("輔助圖表二：全體受試者分佈箱形圖（dm_boxplot_overall.png）", level=2)
    h2.style.font.name = '微軟正黑體'
    
    img_overall_path = "dm_boxplot_overall.png"
    if os.path.exists(img_overall_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_overall_path, width=Inches(6.2))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(12)
        r_cap = p_cap.add_run("圖 2：全體受試者 (N=22) 之 BMI、血糖與收縮壓整體分佈狀況")
        r_cap.font.size = Pt(9.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    # 嵌入圖表 3
    h2 = doc.add_heading("輔助圖表三：組別 × 性別 雙因子箱形圖（dm_boxplot_gender.png）", level=2)
    h2.style.font.name = '微軟正黑體'
    
    img_gender_path = "dm_boxplot_gender.png"
    if os.path.exists(img_gender_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_gender_path, width=Inches(6.2))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(12)
        r_cap = p_cap.add_run("圖 3：依組別與性別次分組之 BMI、血糖與收縮壓交互箱形圖")
        r_cap.font.size = Pt(9.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    # -------------------------------------------------------------
    # 第六章：接手後的實務建議與注意事項
    # -------------------------------------------------------------
    h1 = doc.add_heading("六、 接手後的實務建議與研究局限性", level=1)
    h1.style.font.name = '微軟正黑體'
    h1.style.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p = doc.add_paragraph()
    p.add_run(
        "在撰寫結案報告、簡報或將本數據延伸應用時，有以下幾點實務細節建議特別留意：\n\n"
        "1. "
    )
    p.add_run("樣本數規模提醒（N=22）：").font.bold = True
    p.add_run(
        "目前資料集總人數為 22 人，每組 11 人。雖然組間差異極度鮮明（p < 0.001），"
        "但在醫學推論上應定位為「前導性研究（Pilot Study）」。向外報告時，避免誇大推論為全人口母群體結論。\n\n"
        "2. "
    )
    p.add_run("年齡與性別的潛在干擾（Confounding Factor）：").font.bold = True
    p.add_run(
        "原始資料中包含年齡（Age），分析顯示第 1 組病患平均年齡偏高（平均 51.9 歲），而第 2 組平均為 45.1 歲。"
        "收縮壓的上升有一部分可能受年齡影響。若未來有擴充分析需求，建議可採用「共變數分析（ANCOVA）」或「多元線性迴歸（Multiple Linear Regression）」，"
        "將年齡與性別作為控制變項排除干擾。\n\n"
        "3. "
    )
    p.add_run("中文字型顯示相容性：").font.bold = True
    p.add_run(
        "在 macOS 系統執行繪圖時，程式已預設載入 'Heiti TC' 與 'PingFang TC'；"
        "若未來將腳本移轉至 Windows 系統伺服器執行，請確認作業系統已安裝支援中文的字型（例如微軟正黑體 'Microsoft JhengHei'），"
        "避免 Matplotlib 繪圖時文字變成方塊亂碼。\n\n"
        "4. "
    )
    p.add_run("圖表發表要求：").font.bold = True
    p.add_run(
        "產出的所有圖檔解析度皆已設定為 300 DPI（學術期刊標準印刷畫質）。若期刊要求 TIFF 或 PDF 向量格式，"
        "只需修改 analysis.py 中的 plt.savefig('檔名.pdf') 即可無損輸出。"
    )

    # 結語簽名檔
    p_sign = doc.add_paragraph()
    p_sign.paragraph_format.space_before = Pt(20)
    p_sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_s = p_sign.add_run("交接整理：生醫數據分析工作小組\n檔案產出日期：2026 年 10 月\n")
    r_s.font.size = Pt(10)
    r_s.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    # 儲存檔案
    output_filename = "DM資料分析專案交接與統計方法手冊.docx"
    doc.save(output_filename)
    print(f"Word 文件已成功儲存至: {output_filename}")

if __name__ == '__main__':
    build_handover_document()
